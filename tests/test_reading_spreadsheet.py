import unittest
from io import BytesIO

from openpyxl import load_workbook

from solar_crm.reading_spreadsheet import (
    READING_COLUMNS,
    build_reading_template,
    read_reading_upload,
    spreadsheet_number,
    validate_reading_rows,
)


class NamedBuffer(BytesIO):
    def __init__(self, data: bytes, name: str):
        super().__init__(data)
        self.name = name

    def getvalue(self) -> bytes:
        return super().getvalue()


class ReadingSpreadsheetTests(unittest.TestCase):
    def setUp(self):
        self.plants = [
            {
                "id": 7,
                "client_name": "Cliente teste",
                "name": "Usina teste",
                "expected_monthly_kwh": 4200,
            }
        ]

    def test_excel_template_has_working_sheets_and_plant_row(self):
        content = build_reading_template(self.plants, "2026-09-01")
        workbook = load_workbook(BytesIO(content), data_only=True)

        self.assertEqual(workbook.sheetnames, ["Leituras", "Usinas", "Instruções"])
        self.assertEqual(
            [workbook["Leituras"].cell(1, column).value for column in range(1, 14)],
            READING_COLUMNS,
        )
        self.assertEqual(workbook["Leituras"]["A2"].value, 7)
        self.assertEqual(workbook["Leituras"]["B2"].value.date().isoformat(), "2026-09-01")
        self.assertEqual(workbook["Leituras"].freeze_panes, "A2")
        self.assertGreaterEqual(len(workbook["Leituras"].data_validations.dataValidation), 3)
        self.assertEqual(workbook["Usinas"]["C2"].value, "Usina teste")

    def test_generated_excel_can_be_imported(self):
        content = build_reading_template(self.plants, "2026-09-01")
        workbook = load_workbook(BytesIO(content))
        row = workbook["Leituras"][2]
        row[2].value = 1000
        row[3].value = 3900
        row[7].value = 250
        row[8].value = 900
        output = BytesIO()
        workbook.save(output)

        frame = read_reading_upload(NamedBuffer(output.getvalue(), "modelo_leituras.xlsx"))

        self.assertEqual(len(frame), 1)
        self.assertEqual(frame.iloc[0]["usina_id"], 7)
        self.assertEqual(frame.iloc[0]["geracao_kwh"], 3900)

    def test_old_semicolon_csv_and_brazilian_numbers_are_supported(self):
        csv = (
            "usina_id;mes_referencia;consumo_kwh;geracao_kwh;valor_fatura_rs;custo_sem_solar_rs\n"
            "7;2026-09-01;1.250,5;4.100,8;350,20;990,70\n"
        ).encode("utf-8-sig")
        frame = read_reading_upload(NamedBuffer(csv, "leituras.csv"))

        self.assertEqual(len(frame), 1)
        self.assertEqual(spreadsheet_number(frame.iloc[0]["consumo_kwh"]), 1250.5)
        self.assertEqual(spreadsheet_number(frame.iloc[0]["geracao_kwh"]), 4100.8)

    def test_blank_template_rows_are_not_imported_as_zero_readings(self):
        content = build_reading_template(self.plants, "2026-09-01")
        frame = read_reading_upload(NamedBuffer(content, "modelo_leituras.xlsx"))

        records, errors = validate_reading_rows(frame, {7})

        self.assertEqual(records, [])
        self.assertEqual(errors, ["Nenhuma linha preenchida foi encontrada."])

    def test_complete_rows_are_normalized_before_import(self):
        content = build_reading_template(self.plants, "2026-09-01")
        workbook = load_workbook(BytesIO(content))
        row = workbook["Leituras"][2]
        row[2].value = "1.250,50"
        row[3].value = "4.100,80"
        row[7].value = "350,20"
        row[8].value = "990,70"
        output = BytesIO()
        workbook.save(output)
        frame = read_reading_upload(NamedBuffer(output.getvalue(), "leituras.xlsx"))

        records, errors = validate_reading_rows(frame, {7})

        self.assertEqual(errors, [])
        self.assertEqual(records[0]["consumption_kwh"], 1250.5)
        self.assertEqual(records[0]["generation_kwh"], 4100.8)


if __name__ == "__main__":
    unittest.main()
