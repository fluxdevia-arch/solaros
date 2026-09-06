from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


READING_COLUMNS = [
    "usina_id",
    "mes_referencia",
    "consumo_kwh",
    "geracao_kwh",
    "energia_injetada_kwh",
    "energia_compensada_kwh",
    "tarifa_rs_kwh",
    "valor_fatura_rs",
    "custo_sem_solar_rs",
    "disponibilidade_pct",
    "horas_indisponivel",
    "ocorrencias",
    "observacoes",
]

REQUIRED_READING_COLUMNS = {
    "usina_id",
    "mes_referencia",
    "consumo_kwh",
    "geracao_kwh",
    "valor_fatura_rs",
    "custo_sem_solar_rs",
}

HEADER_FILL = PatternFill("solid", fgColor="173B2F")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
ERROR_FILL = PatternFill("solid", fgColor="F4CCCC")


def build_reading_template(plants: list[dict], reference_month: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Leituras"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:M{max(len(plants) + 1, 2)}"

    for column, header in enumerate(READING_COLUMNS, start=1):
        cell = sheet.cell(1, column, header)
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 34

    for row_number, plant in enumerate(plants, start=2):
        values = [
            int(plant["id"]),
            pd.to_datetime(reference_month).date(),
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            100,
            0,
            0,
            None,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, value)
            if column >= 2:
                cell.fill = INPUT_FILL
            cell.alignment = Alignment(vertical="top", wrap_text=column == 13)
        sheet.cell(row_number, 2).number_format = "yyyy-mm-dd"
        for column in range(3, 12):
            sheet.cell(row_number, column).number_format = "0.00"
        sheet.cell(row_number, 12).number_format = "0"

    last_validation_row = max(len(plants) + 1, 500)
    non_negative = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    availability = DataValidation(type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True)
    incidents = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    sheet.add_data_validation(non_negative)
    sheet.add_data_validation(availability)
    sheet.add_data_validation(incidents)
    non_negative.add(f"C2:I{last_validation_row}")
    non_negative.add(f"K2:K{last_validation_row}")
    availability.add(f"J2:J{last_validation_row}")
    incidents.add(f"L2:L{last_validation_row}")
    sheet.conditional_formatting.add(
        f"J2:J{last_validation_row}",
        CellIsRule(operator="notBetween", formula=["0", "100"], fill=ERROR_FILL),
    )

    widths = [12, 17, 16, 16, 22, 25, 17, 19, 22, 23, 22, 14, 38]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    plants_sheet = workbook.create_sheet("Usinas")
    plants_sheet.sheet_view.showGridLines = False
    plants_sheet.freeze_panes = "A2"
    plant_headers = ["usina_id", "cliente", "usina", "geracao_esperada_kwh_mes"]
    for column, header in enumerate(plant_headers, start=1):
        cell = plants_sheet.cell(1, column, header)
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
    for row_number, plant in enumerate(plants, start=2):
        plants_sheet.append([
            int(plant["id"]),
            plant.get("client_name") or "",
            plant.get("name") or "",
            float(plant.get("expected_monthly_kwh") or 0),
        ])
        plants_sheet.cell(row_number, 4).number_format = "0.00"
    for column, width in enumerate([12, 28, 30, 28], start=1):
        plants_sheet.column_dimensions[get_column_letter(column)].width = width

    instructions = workbook.create_sheet("Instruções")
    instructions.sheet_view.showGridLines = False
    instructions["A1"] = "Como preencher"
    instructions["A1"].font = Font(size=16, bold=True, color="173B2F")
    instructions["A3"] = "1. Preencha somente a aba Leituras."
    instructions["A4"] = "2. Não altere os nomes das colunas da primeira linha."
    instructions["A5"] = "3. Use uma linha por usina e mês. O mês deve estar no formato AAAA-MM-DD."
    instructions["A6"] = "4. Campos amarelos são editáveis. Valores de energia e dinheiro devem ser números."
    instructions["A7"] = "5. A aba Usinas serve para consultar o ID correto de cada instalação."
    instructions["A9"] = "Campos obrigatórios"
    instructions["A9"].font = Font(bold=True)
    instructions["A10"] = ", ".join(sorted(REQUIRED_READING_COLUMNS))
    instructions["A10"].alignment = Alignment(wrap_text=True)
    instructions.column_dimensions["A"].width = 105

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def read_reading_upload(uploaded) -> pd.DataFrame:
    filename = str(getattr(uploaded, "name", "") or "").lower()
    data = uploaded.getvalue() if hasattr(uploaded, "getvalue") else uploaded.read()
    if filename.endswith(".xlsx"):
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        if "Leituras" not in workbook.sheetnames:
            raise ValueError("A planilha precisa conter a aba Leituras.")
        rows = workbook["Leituras"].values
        headers = next(rows, None)
        if not headers:
            return pd.DataFrame(columns=READING_COLUMNS)
        frame = pd.DataFrame(rows, columns=[str(value or "").strip() for value in headers])
    else:
        frame = pd.read_csv(BytesIO(data), sep=None, engine="python", encoding="utf-8-sig")

    frame.columns = [str(column).strip() for column in frame.columns]
    frame = frame.dropna(how="all")
    missing = sorted(REQUIRED_READING_COLUMNS - set(frame.columns))
    if missing:
        raise ValueError("Colunas obrigatórias ausentes: " + ", ".join(missing))
    return frame


def spreadsheet_number(value, default: float = 0) -> float:
    if value is None or pd.isna(value) or str(value).strip() == "":
        return float(default)
    if isinstance(value, str):
        text = value.strip().replace("R$", "").replace("%", "").replace(" ", "")
        if "," in text:
            text = text.replace(".", "").replace(",", ".")
        value = text
    return float(value)


def validate_reading_rows(frame: pd.DataFrame, valid_plant_ids: set[int]) -> tuple[list[dict], list[str]]:
    records: list[dict] = []
    errors: list[str] = []
    value_columns = ["consumo_kwh", "geracao_kwh", "valor_fatura_rs", "custo_sem_solar_rs"]

    def is_blank(value) -> bool:
        return value is None or pd.isna(value) or str(value).strip() == ""

    for row_number, (_, row) in enumerate(frame.iterrows(), start=2):
        if all(is_blank(row.get(column)) for column in value_columns):
            continue
        missing_values = [column for column in value_columns if is_blank(row.get(column))]
        if missing_values:
            errors.append(f"linha {row_number}: preencha " + ", ".join(missing_values))
            continue
        try:
            plant_id = int(spreadsheet_number(row.get("usina_id")))
            if plant_id not in valid_plant_ids:
                raise ValueError(f"usina_id {plant_id} não encontrado")
            month = pd.to_datetime(row.get("mes_referencia"), errors="raise").date().replace(day=1).isoformat()
            record = {
                "plant_id": plant_id,
                "reference_month": month,
                "consumption_kwh": spreadsheet_number(row.get("consumo_kwh")),
                "generation_kwh": spreadsheet_number(row.get("geracao_kwh")),
                "injected_kwh": spreadsheet_number(row.get("energia_injetada_kwh")),
                "compensated_kwh": spreadsheet_number(row.get("energia_compensada_kwh")),
                "tariff": spreadsheet_number(row.get("tarifa_rs_kwh")),
                "billed_amount": spreadsheet_number(row.get("valor_fatura_rs")),
                "reference_amount": spreadsheet_number(row.get("custo_sem_solar_rs")),
                "availability_pct": spreadsheet_number(row.get("disponibilidade_pct"), 100),
                "downtime_hours": spreadsheet_number(row.get("horas_indisponivel")),
                "incidents": int(spreadsheet_number(row.get("ocorrencias"))),
                "failure_notes": "" if is_blank(row.get("observacoes")) else str(row.get("observacoes")),
            }
            numeric_values = [record[column] for column in (
                "consumption_kwh", "generation_kwh", "injected_kwh", "compensated_kwh",
                "tariff", "billed_amount", "reference_amount", "downtime_hours", "incidents",
            )]
            if any(value < 0 for value in numeric_values):
                raise ValueError("os valores numéricos não podem ser negativos")
            if not 0 <= record["availability_pct"] <= 100:
                raise ValueError("disponibilidade_pct deve ficar entre 0 e 100")
            records.append(record)
        except (TypeError, ValueError) as exc:
            errors.append(f"linha {row_number}: {exc}")

    if not records and not errors:
        errors.append("Nenhuma linha preenchida foi encontrada.")
    return records, errors
